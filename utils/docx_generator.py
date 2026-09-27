from docx import Document

def generate_docx(text, filename):

    doc = Document()

    doc.add_heading(
        'LegalEase Document',
        level=1
    )

    doc.add_paragraph(text)

    doc.save(filename)

    return filename